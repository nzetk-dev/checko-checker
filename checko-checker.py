import json
import logging
import os
import sys
import time
import sqlite3
import csv
import hashlib
from datetime import datetime
from typing import Optional, Dict, List, Any
from dataclasses import dataclass, asdict
from pathlib import Path

try:
    import requests
except ImportError:
    print("Необходимо установить библиотеку requests: pip install requests")
    sys.exit(1)

try:
    from cryptography.fernet import Fernet
except ImportError:
    print("Необходимо установить библиотеку cryptography: pip install cryptography")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("Для экспорта в Excel установите: pip install openpyxl")
    openpyxl = None

# Конфигурация
CONFIG_FILE = "config.json"
LOG_FILE = "checko_checker.log"
CACHE_DB = "cache.db"
HISTORY_DB = "history.db"
KEY_FILE = "secret.key"

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(LOG_FILE, encoding='utf-8'),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger(__name__)


# ==================== ШИФРОВАНИЕ ====================

class KeyManager:
    """Управление шифрованием API-ключа"""
    
    def __init__(self):
        self.key_file = Path(KEY_FILE)
        self.key = None
    
    def generate_key(self):
        """Генерация нового ключа шифрования"""
        self.key = Fernet.generate_key()
        with open(self.key_file, 'wb') as f:
            f.write(self.key)
        logger.info("Создан новый ключ шифрования")
    
    def load_key(self):
        """Загрузка существующего ключа"""
        if not self.key_file.exists():
            self.generate_key()
        else:
            with open(self.key_file, 'rb') as f:
                self.key = f.read()
    
    def encrypt(self, data: str) -> str:
        """Шифрование строки"""
        self.load_key()
        f = Fernet(self.key)
        encrypted = f.encrypt(data.encode())
        return encrypted.decode()
    
    def decrypt(self, encrypted_data: str) -> str:
        """Расшифровка строки"""
        self.load_key()
        f = Fernet(self.key)
        decrypted = f.decrypt(encrypted_data.encode())
        return decrypted.decode()


# ==================== КЭШ ====================

class CacheManager:
    """Кэширование запросов для экономии лимитов"""
    
    def __init__(self):
        self.conn = sqlite3.connect(CACHE_DB)
        self._create_table()
    
    def _create_table(self):
        """Создание таблицы кэша"""
        self.conn.execute('''
            CREATE TABLE IF NOT EXISTS cache (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                request_hash TEXT UNIQUE,
                endpoint TEXT,
                response TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        self.conn.commit()
    
    def _get_hash(self, endpoint: str, params: Dict) -> str:
        """Генерация хэша запроса"""
        sorted_params = sorted(params.items())
        request_str = f"{endpoint}:{str(sorted_params)}"
        return hashlib.md5(request_str.encode()).hexdigest()
    
    def get(self, endpoint: str, params: Dict) -> Optional[Dict]:
        """Получение данных из кэша"""
        request_hash = self._get_hash(endpoint, params)
        cursor = self.conn.execute(
            'SELECT response FROM cache WHERE request_hash = ?',
            (request_hash,)
        )
        result = cursor.fetchone()
        
        if result:
            logger.info(f"Данные получены из кэша: {endpoint}")
            return json.loads(result[0])
        return None
    
    def set(self, endpoint: str, params: Dict, response: Dict):
        """Сохранение данных в кэш"""
        request_hash = self._get_hash(endpoint, params)
        self.conn.execute(
            'INSERT OR REPLACE INTO cache (request_hash, endpoint, response) VALUES (?, ?, ?)',
            (request_hash, endpoint, json.dumps(response, ensure_ascii=False))
        )
        self.conn.commit()
    
    def clear_old(self, days: int = 7):
        """Очистка старого кэша"""
        self.conn.execute(
            'DELETE FROM cache WHERE created_at < datetime("now", "-{} days")'.format(days)
        )
        self.conn.commit()
    
    def close(self):
        """Закрытие соединения"""
        self.conn.close()


# ==================== ИСТОРИЯ ====================

class HistoryManager:
    """Хранение истории проверок"""
    
    def __init__(self):
        self.conn = sqlite3.connect(HISTORY_DB)
        self._create_table()
    
    def _create_table(self):
        """Создание таблицы истории"""
        self.conn.execute('''
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                inn TEXT,
                company_name TEXT,
                ogrn TEXT,
                status TEXT,
                director TEXT,
                address TEXT,
                okved TEXT,
                checked_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        self.conn.commit()
    
    def add_record(self, company_data: Dict):
        """Добавление записи в историю"""
        self.conn.execute('''
            INSERT INTO history (inn, company_name, ogrn, status, director, address, okved)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (
            company_data.get('inn', ''),
            company_data.get('name', ''),
            company_data.get('ogrn', ''),
            company_data.get('status', ''),
            company_data.get('director', ''),
            company_data.get('address', ''),
            json.dumps(company_data.get('okved_list', []), ensure_ascii=False)
        ))
        self.conn.commit()
    
    def get_all(self) -> List[Dict]:
        """Получение всей истории"""
        cursor = self.conn.execute(
            'SELECT inn, company_name, ogrn, status, director, address, okved, checked_at FROM history ORDER BY checked_at DESC'
        )
        columns = ['inn', 'company_name', 'ogrn', 'status', 'director', 'address', 'okved', 'checked_at']
        return [dict(zip(columns, row)) for row in cursor.fetchall()]
    
    def close(self):
        """Закрытие соединения"""
        self.conn.close()


# ==================== ЭКСПОРТ ====================

class ExportManager:
    """Экспорт данных в CSV и Excel"""
    
    @staticmethod
    def to_csv(data: List[Dict], filename: str = "export.csv"):
        """Экспорт в CSV"""
        if not data:
            logger.warning("Нет данных для экспорта")
            return
        
        fieldnames = ['inn', 'company_name', 'ogrn', 'status', 'director', 'address', 'okved', 'checked_at']
        
        with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)
        
        logger.info(f"Данные экспортированы в {filename}")
        print(f"✓ Данные экспортированы в {filename}")
    
    @staticmethod
    def to_excel(data: List[Dict], filename: str = "export.xlsx"):
        """Экспорт в Excel"""
        if not data:
            logger.warning("Нет данных для экспорта")
            return
        
        if openpyxl is None:
            print("Библиотека openpyxl не установлена. Экспорт в Excel невозможен.")
            return
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "История проверок"
        
        # Заголовки
        headers = ['ИНН', 'Название', 'ОГРН', 'Статус', 'Директор', 'Адрес', 'ОКВЭД', 'Дата проверки']
        ws.append(headers)
        
        # Стилизация заголовков
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        
        # Данные
        for record in data:
            okved = json.loads(record.get('okved', '[]')) if isinstance(record.get('okved'), str) else record.get('okved', [])
            okved_str = ', '.join(okved) if okved else ''
            
            ws.append([
                record.get('inn', ''),
                record.get('company_name', ''),
                record.get('ogrn', ''),
                record.get('status', ''),
                record.get('director', ''),
                record.get('address', ''),
                okved_str,
                record.get('checked_at', '')
            ])
        
        # Автоширина столбцов
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(filename)
        logger.info(f"Данные экспортированы в {filename}")
        print(f"✓ Данные экспортированы в {filename}")


# ==================== API ====================

@dataclass
class CompanyData:
    """Структура данных компании"""
    inn: str
    ogrn: str
    name: str
    kpp: str
    address: str
    director: str
    okved_list: List[str]
    registration_date: str
    status: str
    contacts: Dict[str, Any]
    finances: Optional[Dict[str, Any]] = None


class CheckoAPI:
    """Класс для работы с API Checko.ru"""
    
    BASE_URL = "https://api.checko.ru/v2"
    
    def __init__(self, api_key: str, cache: CacheManager):
        self.api_key = api_key
        self.cache = cache
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'CheckoChecker/2.0',
            'Accept': 'application/json'
        })
    
    def _make_request(self, endpoint: str, params: Dict[str, Any]) -> Dict[str, Any]:
        """Выполнение запроса с проверкой кэша"""
        # Проверяем кэш
        cached_data = self.cache.get(endpoint, params)
        if cached_data:
            return cached_data
        
        url = f"{self.BASE_URL}/{endpoint}"
        params['key'] = self.api_key
        
        try:
            logger.debug(f"Запрос к {url}")
            response = self.session.get(url, params=params, timeout=30)
            
            # Обработка лимита запросов (32 запроса/сек)
            if response.status_code == 429:
                logger.warning("Превышен лимит запросов. Ожидание 1 секунды...")
                time.sleep(1)
                return self._make_request(endpoint, params)
            
            response.raise_for_status()
            data = response.json()
            
            # Проверка ошибок API
            if 'error' in data:
                logger.error(f"Ошибка API: {data['error']}")
                raise Exception(f"Ошибка API: {data['error']}")
            
            # Сохраняем в кэш
            self.cache.set(endpoint, params, data)
            
            return data
            
        except requests.exceptions.Timeout:
            logger.error("Превышено время ожидания запроса")
            raise
        except requests.exceptions.ConnectionError:
            logger.error("Ошибка подключения к серверу")
            raise
        except requests.exceptions.RequestException as e:
            logger.error(f"Ошибка запроса: {e}")
            raise
    
    def get_company(self, inn: str) -> Optional[CompanyData]:
        """Получение данных организации по ИНН"""
        logger.info(f"Запрос данных компании с ИНН: {inn}")
        
        try:
            # Получение базовой информации
            company_data = self._make_request('company', {'inn': inn})
            
            if not company_data or 'data' not in company_data:
                logger.warning(f"Компания с ИНН {inn} не найдена")
                return None
            
            data = company_data['data']
            
            # Извлечение основных данных
            company = CompanyData(
                inn=data.get('inn', ''),
                ogrn=data.get('ogrn', ''),
                name=data.get('name', ''),
                kpp=data.get('kpp', ''),
                address=data.get('address', {}).get('value', ''),
                director=data.get('manager', {}).get('name', ''),
                okved_list=[okved.get('main', '') for okved in data.get('okved', [])],
                registration_date=data.get('reg_date', ''),
                status=data.get('status', ''),
                contacts=data.get('contacts', {})
            )
            
            # Получение финансовой отчётности
            try:
                finances_data = self._make_request('finances', {'inn': inn})
                if finances_data and 'data' in finances_data:
                    company.finances = finances_data['data']
            except Exception as e:
                logger.warning(f"Не удалось получить финансовые данные: {e}")
            
            return company
            
        except Exception as e:
            logger.error(f"Ошибка при получении данных компании {inn}: {e}")
            return None
    
    def search_company(self, query: str) -> List[Dict[str, Any]]:
        """Поиск компаний по названию"""
        logger.info(f"Поиск компаний по запросу: {query}")
        
        try:
            results = self._make_request('search', {
                'by': 'name',
                'obj': 'org',
                'query': query
            })
            
            if 'data' in results:
                return results['data']
            return []
            
        except Exception as e:
            logger.error(f"Ошибка при поиске: {e}")
            return []
    
    def validate_inn(self, inn: str) -> bool:
        """Валидация ИНН (10 цифр для юрлица)"""
        if not inn.isdigit() or len(inn) != 10:
            return False
        
        # Контрольные числа ИНН
        weights = [2, 4, 10, 3, 5, 9, 4, 6, 8]
        checksum = sum(int(inn[i]) * weights[i] for i in range(9)) % 11 % 10
        
        return checksum == int(inn[9])


# ==================== ОСНОВНОЕ ПРИЛОЖЕНИЕ ====================

class CheckoChecker:
    """Основной класс приложения"""
    
    def __init__(self):
        self.api: Optional[CheckoAPI] = None
        self.cache = CacheManager()
        self.history = HistoryManager()
        self.export = ExportManager()
        self.key_manager = KeyManager()
        self.load_config()
    
    def load_config(self):
        """Загрузка конфигурации"""
        config_path = Path(CONFIG_FILE)
        
        if not config_path.exists():
            logger.info("Конфиг-файл не найден. Создаю новый...")
            api_key = input("Введите ваш API-ключ Checko.ru: ").strip()
            
            # Шифруем ключ
            encrypted_key = self.key_manager.encrypt(api_key)
            
            config = {
                'api_key_encrypted': encrypted_key,
                'default_mode': 'single'
            }
            
            with open(config_path, 'w', encoding='utf-8') as f:
                json.dump(config, f, indent=2, ensure_ascii=False)
            
            logger.info(f"Конфиг сохранён в {CONFIG_FILE}")
        
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
        
        encrypted_key = config.get('api_key_encrypted')
        if not encrypted_key:
            logger.error("API-ключ не настроен. Отредактируйте config.json")
            sys.exit(1)
        
        # Расшифровываем ключ
        try:
            api_key = self.key_manager.decrypt(encrypted_key)
        except Exception as e:
            logger.error(f"Ошибка расшифровки API-ключа: {e}")
            logger.error("Возможно, файл ключа был изменён. Удалите secret.key и config.json для пересоздания.")
            sys.exit(1)
        
        self.api = CheckoAPI(api_key, self.cache)
    
    def print_company_info(self, company: CompanyData):
        """Красивый вывод информации о компании"""
        print("\n" + "="*80)
        print(f"ИНН: {company.inn}")
        print(f"ОГРН: {company.ogrn}")
        print(f"Название: {company.name}")
        print(f"КПП: {company.kpp}")
        print(f"Статус: {company.status}")
        print(f"Дата регистрации: {company.registration_date}")
        print(f"Директор: {company.director}")
        print(f"Адрес: {company.address}")
        
        print(f"\nВиды деятельности (ОКВЭД):")
        for okved in company.okved_list:
            print(f"  - {okved}")
        
        if company.contacts:
            print(f"\nКонтакты:")
            for key, value in company.contacts.items():
                if value:
                    print(f"  {key}: {value}")
        
        if company.finances:
            print(f"\nФинансовая отчётность:")
            for year, data in company.finances.items():
                if isinstance(data, dict):
                    print(f"  {year}:")
                    for key, value in data.items():
                        print(f"    {key}: {value}")
        
        print("="*80 + "\n")
    
    def single_check(self):
        """Одиночная проверка по ИНН"""
        while True:
            inn = input("\nВведите ИНН (10 цифр) или 'q' для выхода: ").strip()
            
            if inn.lower() == 'q':
                break
            
            if not self.api.validate_inn(inn):
                logger.error("Некорректный ИНН. Должно быть 10 цифр.")
                continue
            
            company = self.api.get_company(inn)
            
            if company:
                self.print_company_info(company)
                # Сохраняем в историю
                self.history.add_record(asdict(company))
                logger.info("Результат сохранён в историю")
            else:
                print(f"Компания с ИНН {inn} не найдена или произошла ошибка.")
    
    def search_mode(self):
        """Режим поиска по названию"""
        while True:
            query = input("\nВведите название компании или 'q' для выхода: ").strip()
            
            if query.lower() == 'q':
                break
            
            results = self.api.search_company(query)
            
            if not results:
                print("Ничего не найдено.")
                continue
            
            print(f"\nНайдено компаний: {len(results)}")
            for i, result in enumerate(results[:10], 1):
                print(f"{i}. {result.get('name', 'N/A')} (ИНН: {result.get('inn', 'N/A')})")
            
            choice = input("\nВведите номер для подробной информации или 'q' для выхода: ").strip()
            
            if choice.lower() == 'q':
                break
            
            try:
                idx = int(choice) - 1
                if 0 <= idx < len(results):
                    inn = results[idx].get('inn')
                    if inn:
                        company = self.api.get_company(inn)
                        if company:
                            self.print_company_info(company)
                            # Сохраняем в историю
                            self.history.add_record(asdict(company))
            except ValueError:
                print("Некорректный ввод.")
    
    def batch_check(self):
        """Массовая проверка из файла"""
        file_path = input("Введите путь к файлу со списком ИНН (по одному в строке): ").strip()
        
        if not Path(file_path).exists():
            logger.error(f"Файл {file_path} не найден")
            return
        
        with open(file_path, 'r', encoding='utf-8') as f:
            inns = [line.strip() for line in f if line.strip()]
        
        logger.info(f"Загружено ИНН для проверки: {len(inns)}")
        
        results = []
        output_file = "results.txt"
        
        with open(output_file, 'w', encoding='utf-8') as out:
            out.write("Результаты проверки компаний\n")
            out.write("="*80 + "\n\n")
            
            for i, inn in enumerate(inns, 1):
                logger.info(f"Обработка {i}/{len(inns)}: ИНН {inn}")
                
                if not self.api.validate_inn(inn):
                    out.write(f"ИНН {inn}: Некорректный формат\n\n")
                    continue
                
                company = self.api.get_company(inn)
                
                if company:
                    company_dict = asdict(company)
                    results.append(company_dict)
                    
                    out.write(f"ИНН: {company.inn}\n")
                    out.write(f"Название: {company.name}\n")
                    out.write(f"ОГРН: {company.ogrn}\n")
                    out.write(f"Статус: {company.status}\n")
                    out.write(f"Директор: {company.director}\n")
                    out.write(f"Адрес: {company.address}\n")
                    out.write("-"*80 + "\n\n")
                    
                    # Сохраняем в историю
                    self.history.add_record(company_dict)
                    
                    # Задержка для соблюдения лимитов API
                    time.sleep(0.5)
                else:
                    out.write(f"ИНН {inn}: Не найдена\n\n")
        
        logger.info(f"Результаты сохранены в {output_file}")
        print(f"\nРезультаты сохранены в {output_file}")
    
    def show_history(self):
        """Просмотр истории проверок"""
        history = self.history.get_all()
        
        if not history:
            print("История проверок пуста.")
            return
        
        print(f"\nИстория проверок ({len(history)} записей):")
        print("-"*80)
        
        for i, record in enumerate(history[:20], 1):  # Показываем последние 20
            print(f"{i}. {record['checked_at']} | ИНН: {record['inn']} | {record['company_name']}")
        
        if len(history) > 20:
            print(f"... и ещё {len(history) - 20} записей")
    
    def export_data(self):
        """Экспорт данных"""
        history = self.history.get_all()
        
        if not history:
            print("Нет данных для экспорта.")
            return
        
        print("\nВыберите формат экспорта:")
        print("1. CSV")
        print("2. Excel")
        
        choice = input("Ваш выбор (1-2): ").strip()
        
        if choice == '1':
            self.export.to_csv(history)
        elif choice == '2':
            self.export.to_excel(history)
        else:
            print("Некорректный выбор.")
    
    def run(self):
        """Главный цикл приложения"""
        print("\n" + "="*80)
        print("CHECKO.RU API CHECKER v2.0")
        print("="*80)
        
        # Очищаем старый кэш при запуске
        self.cache.clear_old(days=7)
        
        while True:
            print("\nВыберите режим работы:")
            print("1. Проверка по ИНН (одиночная)")
            print("2. Поиск по названию")
            print("3. Массовая проверка из файла")
            print("4. История проверок")
            print("5. Экспорт данных (CSV/Excel)")
            print("6. Выход")
            
            choice = input("\nВаш выбор (1-6): ").strip()
            
            try:
                if choice == '1':
                    self.single_check()
                elif choice == '2':
                    self.search_mode()
                elif choice == '3':
                    self.batch_check()
                elif choice == '4':
                    self.show_history()
                elif choice == '5':
                    self.export_data()
                elif choice == '6':
                    print("До свидания!")
                    break
                else:
                    print("Некорректный выбор. Введите 1-6.")
            except KeyboardInterrupt:
                print("\n\nПрервано пользователем. До свидания!")
                break
            except Exception as e:
                logger.error(f"Неожиданная ошибка: {e}")
        
        # Закрываем соединения
        self.cache.close()
        self.history.close()


def main():
    """Точка входа"""
    try:
        app = CheckoChecker()
        app.run()
    except Exception as e:
        logger.critical(f"Критическая ошибка: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()